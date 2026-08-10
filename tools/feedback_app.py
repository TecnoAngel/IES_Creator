from __future__ import annotations

import html as html_lib
import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from openpyxl import load_workbook

# De menor a mayor logro. Son caracteres de texto normales (no CSS), así que
# se ven igual pegados en cualquier sitio, aunque el editor de destino elimine
# etiquetas de tabla/estilos (como hace Stilus).
LEVEL_EMOJIS = ("🔴", "🟠", "🟡", "🟢")


def safe_filename(name: str) -> str:
    name = name.strip()
    name = re.sub(r'[\\/:*?"<>|]+', "_", name)
    return name or "Alumno"


def _read_source_bytes(source: Any) -> bytes:
    if isinstance(source, (str, Path)):
        path = Path(source).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"No existe el archivo: {path}")
        return path.read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    raise TypeError("El origen debe ser una ruta o un archivo subido")


def _is_zero(value: Any) -> bool:
    try:
        return float(value) == 0.0
    except (TypeError, ValueError):
        return False


def _esc(value: Any) -> str:
    return html_lib.escape("" if value is None else str(value))


def _set_cell_background(cell, color_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _read_rows(payload: bytes) -> list[tuple]:
    wb = load_workbook(BytesIO(payload), data_only=True)
    sh = wb.active
    return list(sh.iter_rows(values_only=True))


def _parse_feedback_rows(payload: bytes) -> tuple[str, list[str], list[tuple[str, Any, list[Any]]]]:
    """Lee el Excel de feedback y devuelve (etiqueta_actividad, encabezados_items, filas),
    donde cada fila es (alumno, nota, valores)."""
    data = _read_rows(payload)

    if not data or len(data) < 2 or len(data[0]) < 3:
        raise ValueError("La tabla debe tener encabezados y al menos 3 columnas.")

    headers = data[0]
    activity_label = str(headers[0]).strip() if headers[0] is not None else ""
    item_headers = [("" if h is None else str(h).strip()) for h in headers[2:]]

    rows: list[tuple[str, Any, list[Any]]] = []
    for r in data[1:]:
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        student = str(r[0]).strip()
        grade = r[1]
        values = (list(r[2:]) + [None] * len(item_headers))[: len(item_headers)]
        rows.append((student, grade, values))

    if not rows:
        raise ValueError("No se ha encontrado ningún alumno con datos en el Excel.")

    return activity_label, item_headers, rows


def build_student_docx(student: str, activity_label: str, grade: Any, item_headers: list[str], values: list[Any]) -> bytes:
    doc = Document()

    p = doc.add_paragraph()
    run = p.add_run(student)
    run.bold = True

    last_row = 1 + len(item_headers)
    table = doc.add_table(rows=last_row, cols=2)
    table.style = "Table Grid"
    table.autofit = False

    header_row = table.rows[0]
    header_row.cells[0].text = activity_label
    header_row.cells[1].text = "" if grade is None else str(grade)
    for cell in header_row.cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for i, (label, value) in enumerate(zip(item_headers, values), start=1):
        row = table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = "" if value is None else str(value)
        color = "FFC0C0" if _is_zero(value) else "C0FFC0"
        _set_cell_background(row.cells[1], color)

    for row in table.rows:
        row.cells[0].width = Cm(12)
        row.cells[1].width = Cm(2.5)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def run_feedback(source: Any) -> bytes:
    """Genera un ZIP en memoria con un .docx de feedback por alumno."""
    import zipfile

    payload = _read_source_bytes(source)
    activity_label, item_headers, rows = _parse_feedback_rows(payload)

    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for student, grade, values in rows:
            docx_bytes = build_student_docx(student, activity_label, grade, item_headers, values)
            zf.writestr(f"{safe_filename(student)}.docx", docx_bytes)

    return zip_buffer.getvalue()


def _level_emoji(value: Any, max_value: float = 10.0) -> str:
    """🔴🟠🟡🟢 según en qué cuarto de la escala (0-10 por convención) cae el
    valor. Con las notas binarias 0/10 actuales solo se usan 🔴 y 🟢, pero ya
    queda listo para rúbricas con más granularidad (p.ej. 0/3.33/6.67/10)."""
    if value is None:
        return "⚪"
    try:
        ratio = float(value) / max_value if max_value else 0.0
    except (TypeError, ValueError):
        return "⚪"
    if ratio >= 0.75:
        return LEVEL_EMOJIS[3]
    if ratio >= 0.5:
        return LEVEL_EMOJIS[2]
    if ratio >= 0.25:
        return LEVEL_EMOJIS[1]
    return LEVEL_EMOJIS[0]


def _fmt_score(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f"({_esc(value)} sobre 10)"


def build_student_rubric_html(activity_label: str, grade: Any, item_headers: list[str], values: list[Any]) -> str:
    """Rúbrica de un alumno como texto plano con emoji de color por línea
    (🟢🟡🟠🔴), sin tablas ni estilos: en Stilus se ha comprobado que tanto
    <table> como el color por CSS se pierden al pegar, pero el emoji es un
    carácter de texto normal y sobrevive en cualquier sitio. Formato por
    línea: {emoji}{ítem}  (nota sobre 10)."""
    lines = [f"<div><b>{_esc(activity_label)} &mdash; Nota: {_esc(grade)}</b></div>"]
    for label, value in zip(item_headers, values):
        emoji = _level_emoji(value)
        score = _fmt_score(value)
        lines.append(f"<div>{emoji}{_esc(label)}&emsp;{score}</div>")

    return "".join(lines)


def read_students_html(source: Any) -> list[dict]:
    """Igual que run_feedback, pero devuelve por cada alumno su rúbrica en HTML
    en vez de un .docx (para pegar directamente en una plataforma web)."""
    payload = _read_source_bytes(source)
    activity_label, item_headers, rows = _parse_feedback_rows(payload)

    students: list[dict] = []
    for student, grade, values in rows:
        rubric_html = build_student_rubric_html(activity_label, grade, item_headers, values)
        students.append({"name": student, "html": rubric_html})
    return students


_COMPONENT_TEMPLATE = r"""
<div style="font-family:Segoe UI,Arial,sans-serif;color:#222222;color-scheme:light;">
  <style>
    /* Fuerza claro en todo el componente: evita que el modo oscuro del
       navegador reinvierta colores y deje texto invisible sobre el fondo. */
    #fs-toolbar, #fs-editor, #fs-students-grid, #fs-intro-copy-row {
      color-scheme: light;
    }
    #fs-toolbar button {
      border: 1px solid #d7d7e0; background-color: #ffffff; color: #222222;
      border-radius: 6px; padding: 6px 11px; margin-right: 6px; cursor: pointer;
      font-size: 13px;
    }
    #fs-toolbar button:hover { background-color: #f0f0fb; }
    #fs-editor {
      border: 1px solid #d7d7e0; border-radius: 8px; padding: 12px 14px;
      min-height: 110px; margin-top: 8px; font-size: 14px; line-height: 1.5;
      outline: none; background-color: #ffffff; color: #222222;
    }
    #fs-editor:empty:before {
      content: attr(data-placeholder); color: #999999;
    }
    .fs-copy-btn {
      border: none; border-radius: 8px; padding: 10px 14px; cursor: pointer;
      font-size: 13px; font-weight: 600; color: #ffffff;
      background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
      width: 100%; text-align: left; transition: opacity .15s ease;
    }
    .fs-copy-btn:hover { opacity: 0.9; }
    .fs-copy-btn.fs-copied { background: #2e7d32; }
    #fs-students-grid {
      display: grid; grid-template-columns: repeat(auto-fill, minmax(220px, 1fr));
      gap: 10px; margin-top: 10px; max-height: 320px; overflow-y: auto; padding: 4px 2px;
    }
    #fs-intro-copy-row { margin-top: 10px; }
    #fs-intro-copy-row button {
      border: 1px solid #667eea; color: #667eea; background-color: #ffffff;
      border-radius: 6px; padding: 6px 12px; cursor: pointer; font-size: 13px;
    }
    #fs-intro-copy-row button:hover { background-color: #f0f0fb; }
  </style>

  <div id="fs-toolbar">
    <button type="button" onclick="fsExec('bold')"><b>N</b></button>
    <button type="button" onclick="fsExec('italic')"><i>K</i></button>
    <button type="button" onclick="fsExec('underline')"><u>S</u></button>
    <button type="button" onclick="fsExec('insertUnorderedList')">&bull; Lista</button>
  </div>
  <div id="fs-editor" contenteditable="true" data-placeholder="Escribe aquí una introducción general (opcional)..."></div>
  <div id="fs-intro-copy-row">
    <button type="button" id="fs-intro-copy-btn" onclick="fsCopyIntroOnly()">Copiar solo la introducción</button>
    <span id="fs-intro-copy-msg" style="margin-left:8px;font-size:12px;color:#2e7d32;"></span>
  </div>

  <hr style="margin:18px 0;border:none;border-top:1px solid #eee;">

  <div style="font-size:14px;font-weight:600;margin-bottom:4px;">Alumnos (__STUDENT_COUNT__)</div>
  __STUDENTS_BLOCK__

</div>

<script>
  var FS_STUDENTS = __STUDENTS_JSON__;

  function fsExec(cmd) {
    document.getElementById('fs-editor').focus();
    document.execCommand(cmd, false, null);
  }

  function fsCopyRichHtml(htmlContent) {
    var container = document.createElement('div');
    container.innerHTML = htmlContent;
    container.style.position = 'fixed';
    container.style.left = '-9999px';
    container.style.top = '0';
    document.body.appendChild(container);

    var range = document.createRange();
    range.selectNodeContents(container);
    var sel = window.getSelection();
    sel.removeAllRanges();
    sel.addRange(range);

    var ok = false;
    try {
      ok = document.execCommand('copy');
    } catch (err) {
      ok = false;
    }
    sel.removeAllRanges();
    document.body.removeChild(container);
    return ok;
  }

  function fsGetIntroHtml() {
    var editor = document.getElementById('fs-editor');
    if (!editor.innerText || !editor.innerText.trim()) { return ''; }
    return editor.innerHTML;
  }

  function fsCopyIntroOnly() {
    var html = fsGetIntroHtml();
    var msg = document.getElementById('fs-intro-copy-msg');
    if (!html) {
      msg.textContent = 'La introducción está vacía.';
      msg.style.color = '#c62828';
    } else {
      fsCopyRichHtml(html);
      msg.textContent = '¡Copiado!';
      msg.style.color = '#2e7d32';
    }
    setTimeout(function () { msg.textContent = ''; }, 1800);
  }

  function fsCopyForStudent(idx, btnId) {
    var student = FS_STUDENTS[idx];
    var intro = fsGetIntroHtml();
    var combined = intro ? (intro + '<br>' + student.html) : student.html;
    fsCopyRichHtml(combined);

    var btn = document.getElementById(btnId);
    var original = btn.getAttribute('data-original-text');
    btn.textContent = '✓ ¡Copiado!';
    btn.classList.add('fs-copied');
    setTimeout(function () {
      btn.textContent = original;
      btn.classList.remove('fs-copied');
    }, 1500);
  }
</script>
"""


def build_component_html(students: list[dict]) -> str:
    if students:
        buttons = []
        for i, s in enumerate(students):
            btn_id = f"fs-btn-{i}"
            label = html_lib.escape(s["name"])
            buttons.append(
                f'<button type="button" id="{btn_id}" class="fs-copy-btn" '
                f'data-original-text="{label}" onclick="fsCopyForStudent({i}, \'{btn_id}\')">{label}</button>'
            )
        students_block = f'<div id="fs-students-grid">{"".join(buttons)}</div>'
    else:
        students_block = (
            '<div style="color:#888;font-size:13px;margin-top:8px;">'
            "Sube un Excel para ver aquí un botón de copiar por cada alumno.</div>"
        )

    html_out = _COMPONENT_TEMPLATE
    html_out = html_out.replace("__STUDENT_COUNT__", str(len(students)))
    html_out = html_out.replace("__STUDENTS_BLOCK__", students_block)
    html_out = html_out.replace("__STUDENTS_JSON__", json.dumps(students))
    return html_out
