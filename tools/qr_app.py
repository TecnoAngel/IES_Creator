from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

import qrcode
from openpyxl import load_workbook
from PIL import Image
from qrcode.image.styledpil import StyledPilImage
from qrcode.image.styles import colormasks, moduledrawers

ID_COLUMN = "ID"
CODE_COLUMN = "CODE"

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "Qr_Template.xlsx"

# Forma de los módulos (los "píxeles" del QR) y de los 3 ojos de las esquinas.
MODULE_STYLES = {
    "Cuadrado": moduledrawers.SquareModuleDrawer,
    "Cuadros separados": moduledrawers.GappedSquareModuleDrawer,
    "Redondeado": moduledrawers.RoundedModuleDrawer,
    "Puntos": moduledrawers.CircleModuleDrawer,
    "Barras verticales": moduledrawers.VerticalBarsDrawer,
    "Barras horizontales": moduledrawers.HorizontalBarsDrawer,
}
EYE_STYLES = {
    "Cuadrado": moduledrawers.SquareModuleDrawer,
    "Redondeado": moduledrawers.RoundedModuleDrawer,
    "Puntos": moduledrawers.CircleModuleDrawer,
}
FILL_STYLES = [
    "Sólido",
    "Degradado radial",
    "Degradado cuadrado",
    "Degradado horizontal",
    "Degradado vertical",
    "Textura (imagen)",
]

# De más sobrio a más artístico. Cada uno rellena los valores de QRDesign.
PRESETS = {
    "Clásico": dict(module="Cuadrado", eye="Cuadrado", fill="Sólido", color1="#000000", color2="#000000"),
    "Redondeado suave": dict(module="Redondeado", eye="Redondeado", fill="Sólido", color1="#1F2937", color2="#1F2937"),
    "Cuadros separados": dict(module="Cuadros separados", eye="Cuadrado", fill="Sólido", color1="#111111", color2="#111111"),
    "Puntos": dict(module="Puntos", eye="Redondeado", fill="Sólido", color1="#0047AB", color2="#0047AB"),
    "Barras verticales": dict(module="Barras verticales", eye="Redondeado", fill="Degradado vertical", color1="#FF006F", color2="#0047AB"),
    "Barras horizontales": dict(module="Barras horizontales", eye="Redondeado", fill="Degradado horizontal", color1="#667EEA", color2="#764BA2"),
    "Degradado radial": dict(module="Puntos", eye="Redondeado", fill="Degradado radial", color1="#FF006F", color2="#0047AB"),
    "Neón": dict(module="Redondeado", eye="Redondeado", fill="Degradado cuadrado", color1="#7A00FF", color2="#00C2FF"),
    "Textura de imagen": dict(module="Puntos", eye="Redondeado", fill="Textura (imagen)", color1="#000000", color2="#000000"),
}


@dataclass
class QRDesign:
    module: str = "Cuadrado"
    eye: str = "Cuadrado"
    fill: str = "Sólido"
    color1: str = "#000000"  # color principal (o centro/izquierda/arriba del degradado)
    color2: str = "#000000"  # segundo color del degradado
    back: str = "#FFFFFF"
    logo: bytes | None = None  # imagen para el centro
    logo_ratio: float = 0.25  # ancho del logo respecto al QR
    texture: bytes | None = None  # imagen con la que se colorean los módulos

    @classmethod
    def from_preset(cls, name: str, **overrides: Any) -> "QRDesign":
        return cls(**{**PRESETS[name], **overrides})


def safe_filename(name: str) -> str:
    name = name.strip()
    name = re.sub(r'[\\/:*?"<>|]+', "_", name)
    name = re.sub(r"\s+", " ", name)
    return name if name else "qr"


def _read_source_bytes(source: Any) -> bytes:
    if isinstance(source, (str, Path)):
        path = Path(source).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"No existe el archivo: {path}")
        return path.read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    raise TypeError("El origen debe ser una ruta o un archivo subido")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    # Los enteros que openpyxl lee como float (p.ej. 12.0) se escriben sin decimales.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rgb(hex_color: str) -> tuple[int, int, int]:
    h = hex_color.lstrip("#")
    return tuple(int(h[i : i + 2], 16) for i in (0, 2, 4))


def read_qr_rows(payload: bytes) -> list[tuple[str, str]]:
    """Lee la hoja activa (formato de Qr_Template.xlsx) y devuelve (id, code) por fila.
    Las columnas se localizan por su cabecera ID y CODE en la fila 1."""
    wb = load_workbook(BytesIO(payload), data_only=True)
    rows_iter = wb.active.iter_rows(values_only=True)
    try:
        headers = [_cell_text(h).upper() for h in next(rows_iter)]
    except StopIteration:
        raise ValueError("El Excel está vacío.")

    missing = [c for c in (ID_COLUMN, CODE_COLUMN) if c not in headers]
    if missing:
        raise ValueError(f"Faltan columnas en la fila 1: {missing}")
    id_idx = headers.index(ID_COLUMN)
    code_idx = headers.index(CODE_COLUMN)

    rows: list[tuple[str, str]] = []
    for r in rows_iter:
        code = _cell_text(r[code_idx]) if code_idx < len(r) else ""
        if not code:
            continue
        qr_id = _cell_text(r[id_idx]) if id_idx < len(r) else ""
        rows.append((qr_id or code, code))

    if not rows:
        raise ValueError("No se ha encontrado ninguna fila con CODE en el Excel.")
    return rows


def _color_mask(design: QRDesign) -> colormasks.QRColorMask:
    back, c1, c2 = _rgb(design.back), _rgb(design.color1), _rgb(design.color2)
    if design.fill == "Degradado radial":
        return colormasks.RadialGradiantColorMask(back_color=back, center_color=c1, edge_color=c2)
    if design.fill == "Degradado cuadrado":
        return colormasks.SquareGradiantColorMask(back_color=back, center_color=c1, edge_color=c2)
    if design.fill == "Degradado horizontal":
        return colormasks.HorizontalGradiantColorMask(back_color=back, left_color=c1, right_color=c2)
    if design.fill == "Degradado vertical":
        return colormasks.VerticalGradiantColorMask(back_color=back, top_color=c1, bottom_color=c2)
    if design.fill == "Textura (imagen)":
        if not design.texture:
            raise ValueError("El relleno 'Textura (imagen)' necesita que subas una imagen de textura.")
        texture = Image.open(BytesIO(design.texture)).convert("RGB")
        return colormasks.ImageColorMask(back_color=back, color_mask_image=texture)
    return colormasks.SolidFillColorMask(back_color=back, front_color=c1)


def _prepare_logo(logo_bytes: bytes, back_color: str) -> Image.Image:
    """El QR pega el logo estirado a un cuadrado: se centra sobre un lienzo cuadrado
    del color de fondo, con un pequeño margen para separarlo de los módulos."""
    logo = Image.open(BytesIO(logo_bytes)).convert("RGBA")
    side = int(max(logo.size) * 1.15)
    canvas = Image.new("RGBA", (side, side), (*_rgb(back_color), 255))
    canvas.alpha_composite(logo, ((side - logo.width) // 2, (side - logo.height) // 2))
    return canvas


def create_qr_png_bytes(payload: str, design: QRDesign | None = None) -> bytes:
    if design is None:
        design = QRDesign()

    # Con imagen central se tapa parte del QR: corrección H (hasta ~30 % recuperable).
    error_correction = qrcode.constants.ERROR_CORRECT_H if design.logo else qrcode.constants.ERROR_CORRECT_M
    qr = qrcode.QRCode(error_correction=error_correction, box_size=10, border=4)
    qr.add_data(payload)
    qr.make(fit=True)

    kwargs: dict[str, Any] = {
        "image_factory": StyledPilImage,
        "module_drawer": MODULE_STYLES[design.module](),
        "eye_drawer": EYE_STYLES[design.eye](),
        "color_mask": _color_mask(design),
    }
    if design.logo:
        kwargs["embedded_image"] = _prepare_logo(design.logo, design.back)
        kwargs["embedded_image_ratio"] = design.logo_ratio

    img = qr.make_image(**kwargs)
    buffer = BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def run_qr_generation(source: Any, design: QRDesign | None = None) -> bytes:
    """Genera un ZIP en memoria con un PNG por fila del Excel, llamado como su ID."""
    payload = _read_source_bytes(source)
    rows = read_qr_rows(payload)

    used_names: dict[str, int] = {}
    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for qr_id, code in rows:
            # IDs repetidos en el Excel no deben pisarse dentro del ZIP.
            filename = safe_filename(qr_id)
            count = used_names.get(filename, 0)
            used_names[filename] = count + 1
            if count:
                filename = f"{filename}_{count + 1}"
            zf.writestr(f"{filename}.png", create_qr_png_bytes(code, design))

    return zip_buffer.getvalue()
