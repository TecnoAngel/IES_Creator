from __future__ import annotations

import json
import random
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import qrcode
from docx import Document
from docx.shared import Cm
from docxtpl import DocxTemplate
from openpyxl import load_workbook

# Mismo mapeo que encode_answers_for_qr, invertido.
_QR_LETTER_BY_SYMBOL = {"*": "A", "?": "B", "0": "C", "!": "D"}


class NextValue:
    def __init__(self, values: list):
        self._it = iter(values)

    def __str__(self) -> str:
        try:
            v = next(self._it)
        except StopIteration:
            v = ""
        return "" if v is None else str(v)

    def __html__(self) -> str:
        return str(self)


def safe_filename(name: str) -> str:
    name = name.strip()
    name = re.sub(r'[\\/:*?"<>|]+', "_", name)
    name = re.sub(r"\s+", " ", name)
    return name if name else "alumno"


def _read_source_bytes(source: Any) -> bytes:
    if isinstance(source, (str, Path)):
        path = Path(source).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"No existe el archivo: {path}")
        return path.read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    raise TypeError("El origen debe ser una ruta o un archivo subido")


def ensure_template(template_source: Any) -> bytes:
    if template_source is not None:
        return _read_source_bytes(template_source)

    default_template = Path(__file__).resolve().parent.parent / "templates" / "test_template.docx"
    default_template.parent.mkdir(parents=True, exist_ok=True)

    if not default_template.exists():
        doc = Document()
        doc.add_heading("Plantilla de prueba", level=1)
        doc.add_paragraph("Pregunta 1")
        doc.add_paragraph("A. Opción A")
        doc.add_paragraph("B. Opción B")
        doc.add_paragraph("C. Opción C")
        doc.add_paragraph("D. Opción D")
        doc.save(str(default_template))

    return default_template.read_bytes()


def read_usedrange(payload: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(payload), data_only=True)
    sh = wb.active
    rows_iter = sh.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    rows: list[dict] = []
    for row in rows_iter:
        if row is None or all(c in (None, "") for c in row):
            continue
        rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return rows


def _norm(x) -> str:
    if x is None:
        return ""
    return str(x).strip()


def build_shuffled_version(rows: list[dict], rng: random.Random) -> tuple[list[dict], list[str]]:
    rows_copy = [dict(r) for r in rows]
    rng.shuffle(rows_copy)
    out: list[dict] = []
    correct_letters: list[str] = []
    letters = ["A", "B", "C", "D"]

    for r in rows_copy:
        answers = [r.get("ans1", ""), r.get("ans2", ""), r.get("ans3", ""), r.get("ans4", "")]
        correct_value = r.get("ans_correct", "")
        rng.shuffle(answers)
        idx = None
        cv = _norm(correct_value)
        for i, a in enumerate(answers):
            if _norm(a) == cv:
                idx = i
                break
        correct_letter = letters[idx] if idx is not None else "?"
        new_r = dict(r)
        new_r["ans1"], new_r["ans2"], new_r["ans3"], new_r["ans4"] = answers
        out.append(new_r)
        correct_letters.append(correct_letter)
    return out, correct_letters


def encode_answers_for_qr(correct_letters: list[str]) -> str:
    mapping = {"A": "*", "B": "?", "C": "0", "D": "!"}
    encoded = "".join(mapping.get(ch, "?") for ch in correct_letters)
    return "ANS" + encoded


def create_qr_png_bytes(payload: str) -> bytes:
    img = qrcode.make(payload)
    buffer = BytesIO()
    img.save(buffer, format="PNG")
    return buffer.getvalue()


def insert_qr_at_end(docx_bytes: bytes, qr_png_bytes: bytes, width_cm: float = 1.5) -> bytes:
    doc = Document(BytesIO(docx_bytes))
    doc.add_paragraph("")
    doc.add_picture(BytesIO(qr_png_bytes), width=Cm(width_cm))
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_doc(template_bytes: bytes, rows_shuffled: list[dict]) -> bytes:
    questions = [r.get("question", "") for r in rows_shuffled]
    n = len(rows_shuffled)
    ctx = {
        "test": questions,
        "numpreg": list(range(1, n + 1)),
        "ans1": NextValue([r.get("ans1", "") for r in rows_shuffled]),
        "ans2": NextValue([r.get("ans2", "") for r in rows_shuffled]),
        "ans3": NextValue([r.get("ans3", "") for r in rows_shuffled]),
        "ans4": NextValue([r.get("ans4", "") for r in rows_shuffled]),
    }
    doc = DocxTemplate(BytesIO(template_bytes))
    doc.render(ctx)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def run_test_generation(source: Any, template_source: Any = None, num_students: int = 5) -> bytes:
    """Genera un ZIP en memoria con un .docx + QR por alumno."""
    payload = _read_source_bytes(source)
    template_bytes = ensure_template(template_source)

    rows = read_usedrange(payload)
    if not rows:
        raise ValueError("No se han leído filas del Excel.")

    required = {"question", "ans1", "ans2", "ans3", "ans4", "ans_correct"}
    missing = required - set(rows[0].keys())
    if missing:
        raise ValueError(f"Faltan columnas: {sorted(missing)}")

    alumnos = [f"alumno{i}" for i in range(1, num_students + 1)]

    zip_buffer = BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for student in alumnos:
            seed = hash(student) & 0xFFFFFFFF
            rng = random.Random(seed)
            rows_shuffled, correct_letters = build_shuffled_version(rows, rng)
            qr_payload = encode_answers_for_qr(correct_letters)

            docx_bytes = generate_doc(template_bytes, rows_shuffled)
            qr_png_bytes = create_qr_png_bytes(qr_payload)
            docx_bytes = insert_qr_at_end(docx_bytes, qr_png_bytes, width_cm=1.5)

            filename = safe_filename(student)
            zf.writestr(f"{filename}.docx", docx_bytes)
            zf.writestr(f"{filename}_QR.png", qr_png_bytes)

    return zip_buffer.getvalue()


def decode_qr_image(image_bytes: bytes) -> str:
    """Lee el contenido de texto de un QR a partir de una foto/imagen (jpg, png...)."""
    arr = np.frombuffer(image_bytes, dtype=np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError("No se ha podido leer la imagen.")

    detector = cv2.QRCodeDetector()
    data, _points, _straight = detector.detectAndDecode(img)
    if not data:
        raise ValueError("No se ha detectado ningún código QR en la imagen.")
    return data


def decode_answers_from_qr_payload(payload: str) -> list[str]:
    """Invierte encode_answers_for_qr: 'ANS*?0!' -> ['A', 'B', 'C', 'D']."""
    if not payload.startswith("ANS"):
        raise ValueError("El código leído no tiene el formato esperado (debe empezar por 'ANS').")

    letters: list[str] = []
    for symbol in payload[3:]:
        letters.append(_QR_LETTER_BY_SYMBOL.get(symbol, "¿?"))
    if not letters:
        raise ValueError("El código QR no contiene ninguna respuesta.")
    return letters


_TTS_COMPONENT_TEMPLATE = r"""
<div style="font-family:Segoe UI,Arial,sans-serif;color:#222222;color-scheme:light;">
  <style>
    #tts-controls { display: flex; align-items: center; gap: 10px; flex-wrap: wrap; }
    #tts-controls button {
      border: none; border-radius: 8px; padding: 10px 16px; cursor: pointer;
      font-size: 14px; font-weight: 600; color: #ffffff;
      background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    }
    #tts-controls button:hover { opacity: 0.9; }
    #tts-controls button#tts-stop { background: #888888; }
    #tts-rate-row { display: flex; align-items: center; gap: 8px; font-size: 13px; }
    #tts-status {
      margin-top: 12px; font-size: 14px; font-weight: 600; color: #333333;
      min-height: 20px;
    }
    #tts-list { margin-top: 10px; font-size: 13px; color: #555555; line-height: 1.6; }
  </style>

  <div id="tts-controls">
    <button type="button" id="tts-play">▶ Reproducir</button>
    <button type="button" id="tts-stop">■ Detener</button>
    <div id="tts-rate-row">
      <span>Ritmo:</span>
      <input type="range" id="tts-rate" min="0.5" max="2" step="0.1" value="1">
      <span id="tts-rate-label">1.0x</span>
    </div>
  </div>
  <div id="tts-status"></div>
  <div id="tts-list">__PHRASES_HTML__</div>
</div>

<script>
  var TTS_PHRASES = __PHRASES_JSON__;
  var ttsIndex = 0;
  var ttsPlaying = false;

  var rateInput = document.getElementById('tts-rate');
  var rateLabel = document.getElementById('tts-rate-label');
  rateInput.addEventListener('input', function () {
    rateLabel.textContent = parseFloat(rateInput.value).toFixed(1) + 'x';
  });

  function ttsSpeakNext() {
    if (!ttsPlaying || ttsIndex >= TTS_PHRASES.length) {
      if (ttsIndex >= TTS_PHRASES.length) {
        document.getElementById('tts-status').textContent = 'Terminado.';
      }
      ttsPlaying = false;
      return;
    }
    var phrase = TTS_PHRASES[ttsIndex];
    document.getElementById('tts-status').textContent = phrase;
    var utter = new SpeechSynthesisUtterance(phrase);
    utter.lang = 'es-ES';
    utter.rate = parseFloat(rateInput.value);
    utter.onend = function () {
      ttsIndex++;
      ttsSpeakNext();
    };
    utter.onerror = function () {
      ttsIndex++;
      ttsSpeakNext();
    };
    window.speechSynthesis.speak(utter);
  }

  document.getElementById('tts-play').addEventListener('click', function () {
    window.speechSynthesis.cancel();
    ttsIndex = 0;
    ttsPlaying = true;
    ttsSpeakNext();
  });

  document.getElementById('tts-stop').addEventListener('click', function () {
    ttsPlaying = false;
    window.speechSynthesis.cancel();
    document.getElementById('tts-status').textContent = 'Detenido.';
  });
</script>
"""


def build_tts_component_html(letters: list[str]) -> str:
    """Componente de lectura en voz alta (Web Speech API del navegador) de las
    respuestas decodificadas del QR, con control de ritmo."""
    phrases = [f"Respuesta número {i}: {letter}" for i, letter in enumerate(letters, start=1)]
    phrases_html = " &middot; ".join(phrases)

    html_out = _TTS_COMPONENT_TEMPLATE
    html_out = html_out.replace("__PHRASES_JSON__", json.dumps(phrases))
    html_out = html_out.replace("__PHRASES_HTML__", phrases_html)
    return html_out
